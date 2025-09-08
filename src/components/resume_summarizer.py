import ollama
import pdfplumber
from docx import Document
import re
import time
from typing import Optional, List, Dict, Any

from .summary_result import SummaryResult


class ResumeSummarizer:
    
        
    def __init__(self, settings):
        model_settings = settings.get('model_settings', {})
        summary_settings = settings.get('summary_settings', {})
        ollama_settings = settings.get('ollama_settings', {})
        
        self.temperature = model_settings.get('temperature', 0.2)
        self.max_tokens = model_settings.get('max_tokens', 600)
        self.chunk_size = model_settings.get('chunk_size', 3000)
        self.min_summary_length = summary_settings.get('min_length', 430)
        self.max_summary_length = summary_settings.get('max_length', 500)
        
        # Ollama configuration
        self.ollama_base_url = ollama_settings.get('base_url', 'http://localhost:11434')
        self.ollama_timeout = ollama_settings.get('timeout', 60)
        
        # Store settings for prompt generation
        self.settings = settings
        
        # Configure Ollama client
        self._configure_ollama_client()
        
        # Select available model
        self.model = self._select_best_model()
    
    def _configure_ollama_client(self):
        """Configure Ollama client with custom endpoint if specified."""
        if hasattr(self, 'ollama_base_url') and self.ollama_base_url != 'http://localhost:11434':
            import os
            os.environ['OLLAMA_HOST'] = self.ollama_base_url
    
    def _select_best_model(self) -> str:
        try:
            models = ollama.list()
            if models and hasattr(models, 'models'):
                available_models = [model.model for model in models.models]
                
                # Return first available model
                if available_models:
                    return available_models[0]
        except Exception:
            pass
        
        # No models available
        raise RuntimeError("No Ollama models available. Please run: ollama pull <model>")
    
    def extract_text_from_pdf(self, file_path: str) -> str:
        text = ""
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception as e:
            raise RuntimeError(f"Error processing PDF: {str(e)}")
        return text
    
    def extract_text_from_docx(self, file_path: str) -> str:
        try:
            doc = Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text
        except Exception as e:
            raise RuntimeError(f"Error processing DOCX: {str(e)}")
    
    def preprocess_text(self, text: str) -> str:
        # Remove extra whitespace
        text = re.sub(r'\s+', ' ', text)
        # Remove empty lines
        text = re.sub(r'\n\s*\n', '\n', text)
        # Remove special characters that might cause issues
        text = re.sub(r'[^\w\s\.\,\!\?\;\:\-\(\)\[\]\{\}\@\#\$\%\^\&\*\+\=\~\`\'\"\/\\\<\>]', '', text)
        return text.strip()
    
    def chunk_text(self, text: str) -> List[str]:
        words = text.split()
        chunks = []
        current_chunk = []
        current_length = 0
        
        for word in words:
            current_chunk.append(word)
            current_length += len(word) + 1  # +1 for space
            
            if current_length >= self.chunk_size:
                chunks.append(' '.join(current_chunk))
                current_chunk = []
                current_length = 0
        
        if current_chunk:
            chunks.append(' '.join(current_chunk))
        
        return chunks
    
        
        
    def _ensure_summary_length(self, summary: str) -> str:
        if len(summary) < self.min_summary_length:
            # If too short, pad with a standard ending
            padding = " This candidate shows potential for the right role and should be considered for an interview to assess cultural fit and specific capabilities."
            summary = summary + padding
        
        elif len(summary) > self.max_summary_length:
            # If too long, truncate to last sentence within limit
            truncated = summary[:self.max_summary_length]
            # Find the last sentence end
            last_period = truncated.rfind('.')
            last_exclamation = truncated.rfind('!')
            last_question = truncated.rfind('?')
            
            last_sentence_end = max(last_period, last_exclamation, last_question)
            
            if last_sentence_end > 400:  # Only truncate if we have a substantial portion
                summary = truncated[:last_sentence_end + 1]
            else:
                summary = truncated[:480] + "..."
        
        return summary.strip()
    
    def _generate_summary_with_fallback(self, text: str, prompt: str) -> SummaryResult:
        start_time = time.time()
        last_error = None
        
        try:
            response = ollama.chat(
                model=self.model,
                messages=[{"role": "user", "content": prompt}],
                options={
                    "temperature": self.temperature,
                    "num_predict": self.max_tokens
                }
            )
            
            summary = response['message']['content']
            summary = self._ensure_summary_length(summary)
            processing_time = time.time() - start_time
            
            return SummaryResult(
                summary=summary,
                model_used=self.model,
                processing_time=processing_time,
                success=True
            )
            
        except Exception as e:
            last_error = e
        
        # Model failed
        return SummaryResult(
            summary="",
            model_used="",
            processing_time=time.time() - start_time,
            success=False,
            error=str(last_error)
        )
    
    def generate_summary(self, text: str, custom_prompt: str = None) -> SummaryResult:
        try:
            # Use custom prompt if provided, otherwise use prompt from settings
            prompt = custom_prompt or self.settings.get('prompt', '')
            
            if not prompt:
                return SummaryResult(
                    summary="",
                    model_used="",
                    processing_time=0.0,
                    success=False,
                    error="No prompt configured in settings"
                )
            
            # Replace placeholders with actual text and length settings
            prompt = prompt.replace("{document_text}", text)
            prompt = prompt.replace("{min_length}", str(self.min_summary_length))
            prompt = prompt.replace("{max_length}", str(self.max_summary_length))
            
            return self._generate_summary_with_fallback(text, prompt)
            
        except Exception as e:
            return SummaryResult(
                summary="",
                model_used="",
                processing_time=0.0,
                success=False,
                error=str(e)
            )
    
    def process_document(self, file_path: str, file_type: str, 
                        progress_callback=None, custom_prompt: str = None) -> SummaryResult:
        # Extract text based on file type
        if file_type == "pdf":
            text = self.extract_text_from_pdf(file_path)
        elif file_type == "docx":
            text = self.extract_text_from_docx(file_path)
        else:
            return SummaryResult(
                summary="Unsupported file type",
                model_used="",
                processing_time=0.0,
                success=False,
                error="Unsupported file type"
            )
        
        if not text:
            return SummaryResult(
                summary="No text could be extracted from the document",
                model_used="",
                processing_time=0.0,
                success=False,
                error="No text extracted"
            )
        
        # Preprocess text
        text = self.preprocess_text(text)
        
        # Check if text needs to be chunked
        if len(text) > self.chunk_size:
            chunks = self.chunk_text(text)
            summaries = []
            
            for i, chunk in enumerate(chunks):
                if progress_callback:
                    progress_callback(i + 1, len(chunks), f"Processing chunk {i+1}/{len(chunks)}")
                
                chunk_result = self.generate_summary(chunk, custom_prompt)
                if chunk_result.success:
                    summaries.append(chunk_result.summary)
                else:
                    # If any chunk fails, return the error
                    return chunk_result
            
            # Combine summaries
            combined_summary = "\n\n".join(summaries)
            
            # Generate final summary of summaries
            return self.generate_summary(combined_summary, custom_prompt)
        else:
            # Single chunk processing
            return self.generate_summary(text, custom_prompt)
    
    def get_available_models(self) -> List[str]:
        try:
            models = ollama.list()
            if models and hasattr(models, 'models'):
                return [model.model for model in models.models]
        except Exception:
            pass
        return []
    
    